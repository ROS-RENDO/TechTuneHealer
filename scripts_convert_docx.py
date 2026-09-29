import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_styled_paragraph(doc, text, style='Normal', space_after=6, space_before=0, line_spacing=1.15):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = line_spacing
    
    # Simple markdown formatting for bold and italics
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if token.startswith('**') and token.endswith('**'):
            run = p.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('*') and token.endswith('*') and len(token) > 2:
            run = p.add_run(token[1:-1])
            run.italic = True
        elif token.startswith('`') and token.endswith('`'):
            run = p.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(199, 37, 78)
        else:
            p.add_run(token)
    return p

def convert_markdown_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = docx.Document()
    
    # Page Margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base styling
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(33, 37, 41)

    i = 0
    in_code_block = False
    code_block_lines = []
    
    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        
        # Check code block fences
        if line.strip().startswith('```'):
            if in_code_block:
                # Flush code block
                code_text = '\n'.join(code_block_lines)
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                set_cell_background(cell, "F8F9FA")
                set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(code_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(9.0)
                run.font.color.rgb = RGBColor(40, 44, 52)
                code_block_lines = []
                in_code_block = False
            else:
                in_code_block = True
                code_block_lines = []
            i += 1
            continue

        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue

        stripped = line.strip()
        
        if not stripped or stripped == '<br>' or stripped == '---':
            i += 1
            continue

        # Image check: ![caption](./images/...)
        img_match = re.match(r'!\[(.*?)\]\((.*?)\)', stripped)
        if img_match:
            caption = img_match.group(1)
            img_rel_path = img_match.group(2)
            # Resolve image path relative to md_path
            img_abs_path = os.path.normpath(os.path.join(os.path.dirname(md_path), img_rel_path))
            if os.path.exists(img_abs_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(12)
                p_img.paragraph_format.space_after = Pt(4)
                p_img.add_run().add_picture(img_abs_path, width=Inches(6.0))
                
                # Caption
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_after = Pt(12)
                run_cap = p_cap.add_run(caption)
                run_cap.italic = True
                run_cap.font.size = Pt(9.5)
                run_cap.font.color.rgb = RGBColor(108, 117, 125)
            i += 1
            continue

        # Table detection
        if stripped.startswith('|') and stripped.endswith('|'):
            # Collect all table lines
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            # Filter out separator row (e.g., | :--- | :--- |)
            parsed_rows = []
            for t_line in table_lines:
                cells = [c.strip() for c in t_line[1:-1].split('|')]
                # Check if it's separator
                if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                    continue
                parsed_rows.append(cells)
            
            if parsed_rows:
                num_cols = max(len(r) for r in parsed_rows)
                tbl = doc.add_table(rows=len(parsed_rows), cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                
                for r_idx, row_data in enumerate(parsed_rows):
                    for c_idx in range(num_cols):
                        c_text = row_data[c_idx] if c_idx < len(row_data) else ""
                        cell = tbl.cell(r_idx, c_idx)
                        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                        
                        if r_idx == 0:
                            set_cell_background(cell, "2563EB")
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(4)
                            p.paragraph_format.space_after = Pt(4)
                            run = p.add_run(c_text)
                            run.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                            run.font.size = Pt(10)
                        else:
                            if r_idx % 2 == 1:
                                set_cell_background(cell, "F8FAFC")
                            else:
                                set_cell_background(cell, "FFFFFF")
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(3)
                            p.paragraph_format.space_after = Pt(3)
                            
                            # Clean HTML breaks if any
                            clean_text = c_text.replace('<br>', '\n')
                            tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', clean_text)
                            for token in tokens:
                                if token.startswith('**') and token.endswith('**'):
                                    run = p.add_run(token[2:-2])
                                    run.bold = True
                                    run.font.size = Pt(9.5)
                                elif token.startswith('*') and token.endswith('*') and len(token) > 2:
                                    run = p.add_run(token[1:-1])
                                    run.italic = True
                                    run.font.size = Pt(9.5)
                                elif token.startswith('`') and token.endswith('`'):
                                    run = p.add_run(token[1:-1])
                                    run.font.name = 'Consolas'
                                    run.font.size = Pt(9.0)
                                else:
                                    run = p.add_run(token)
                                    run.font.size = Pt(9.5)

                p_after = doc.add_paragraph()
                p_after.paragraph_format.space_after = Pt(6)
            continue

        # Headings
        if stripped.startswith('# '):
            h_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.bold = True
            run.font.size = Pt(18)
            run.font.name = 'Arial'
            run.font.color.rgb = RGBColor(30, 58, 138) # Dark Tech Blue
            i += 1
            continue

        if stripped.startswith('## '):
            h_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.bold = True
            run.font.size = Pt(14)
            run.font.name = 'Arial'
            run.font.color.rgb = RGBColor(37, 99, 235) # Tech Blue
            i += 1
            continue

        if stripped.startswith('### '):
            h_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.bold = True
            run.font.size = Pt(12)
            run.font.name = 'Arial'
            run.font.color.rgb = RGBColor(30, 41, 59)
            i += 1
            continue

        if stripped.startswith('#### '):
            h_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.bold = True
            run.font.size = Pt(11)
            run.font.name = 'Arial'
            i += 1
            continue

        # Bullet list
        if stripped.startswith('- ') or stripped.startswith('* '):
            bullet_text = stripped[2:].strip()
            add_styled_paragraph(doc, bullet_text, style='List Bullet', space_after=3, line_spacing=1.15)
            i += 1
            continue

        # Blockquote / Notes
        if stripped.startswith('> '):
            quote_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.space_before = Pt(4)
            run = p.add_run(quote_text)
            run.italic = True
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(71, 85, 105)
            i += 1
            continue

        # Regular paragraph
        add_styled_paragraph(doc, stripped, style='Normal', space_after=6, line_spacing=1.15)
        i += 1

    doc.save(docx_path)
    print(f"Successfully generated: {docx_path}")

if __name__ == '__main__':
    convert_markdown_to_docx('report.md', 'TechTune_Healer_End_of_Term_Progress_Report.docx')
